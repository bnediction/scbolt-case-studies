#!/usr/bin/env python3
"""Export the integrated annotation results used in the APL case study."""

from __future__ import annotations

import argparse
import csv
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


APL_DIR = Path(__file__).resolve().parents[1]
DEFAULT_PROJECT_DIR = APL_DIR / "project_gsm"
DEFAULT_OUTPUT_DIR = APL_DIR / "results" / "annotation"
ANNOTATION_METADATA = Path("omics/annot/integrated/annot.scbolt.json")
DEA_RESULT = Path("dea/markers.csv")
GOEA_RESULT = Path("goea/basic.xlsx")
SCORING_RESULT = Path("scoring/signature_pvals_adj.csv")
INTEGRATED_RESULTS = {
    DEA_RESULT: Path("omics/dea/integrated/markers.csv"),
    GOEA_RESULT: Path("omics/goea/integrated/basic.xlsx"),
    SCORING_RESULT: Path("omics/scoring/integrated/signature_pvals_adj.csv"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Export the integrated DEA, GOEA and signature-scoring results."
    )
    parser.add_argument(
        "--project-dir",
        type=Path,
        default=DEFAULT_PROJECT_DIR,
        help=f"scBOLT project directory (default: {DEFAULT_PROJECT_DIR})",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_OUTPUT_DIR,
        help=f"annotation result directory (default: {DEFAULT_OUTPUT_DIR})",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="replace existing exported files",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="validate and display the export without copying files",
    )
    return parser.parse_args()


def require_file(path: Path) -> Path:
    if not path.is_file():
        raise FileNotFoundError(f"required integrated result not found: {path}")
    return path


def load_cluster_labels(project_dir: Path) -> dict[str, str]:
    metadata_file = require_file(project_dir / ANNOTATION_METADATA)
    metadata = json.loads(metadata_file.read_text())
    label_value = metadata.get("sensitive_parameters", {}).get("LABEL")
    if not isinstance(label_value, str):
        raise RuntimeError(f"LABEL not found in annotation metadata: {metadata_file}")

    labels = label_value.split()
    if not labels or len(labels) != len(set(labels)):
        raise RuntimeError(
            f"invalid LABEL mapping in annotation metadata: {label_value}"
        )
    return {str(cluster): label for cluster, label in enumerate(labels)}


def build_export_plan(project_dir: Path) -> dict[Path, Path]:
    project_dir = project_dir.resolve()
    return {
        destination: require_file(project_dir / source)
        for destination, source in INTEGRATED_RESULTS.items()
    }


def replace_labels(
    values: list[str], labels: dict[str, str], context: str
) -> list[str]:
    unknown = sorted(set(values) - labels.keys())
    if unknown:
        raise RuntimeError(
            f"unknown cluster identifiers in {context}: {', '.join(unknown)}"
        )
    return [labels[value] for value in values]


def export_dea(source: Path, destination: Path, labels: dict[str, str]) -> None:
    with source.open(newline="", encoding="utf-8") as infile:
        rows = list(csv.reader(infile))
    if not rows or not rows[0] or rows[0][0] != "group":
        raise RuntimeError(f"unexpected DEA table format: {source}")

    cluster_ids = [row[0] for row in rows[1:]]
    for row, label in zip(rows[1:], replace_labels(cluster_ids, labels, "DEA table")):
        row[0] = label

    with destination.open("w", newline="", encoding="utf-8") as outfile:
        csv.writer(outfile, lineterminator="\n").writerows(rows)


def export_goea(source: Path, destination: Path, labels: dict[str, str]) -> None:
    shutil.copy2(source, destination)
    workbook = load_workbook(destination)
    try:
        sheet_names = workbook.sheetnames
        renamed = replace_labels(sheet_names, labels, "GOEA workbook")
        for worksheet, label in zip(workbook.worksheets, renamed):
            worksheet.title = label
        workbook.save(destination)
    finally:
        workbook.close()


def export_scoring(source: Path, destination: Path, labels: dict[str, str]) -> None:
    with source.open(newline="", encoding="utf-8") as infile:
        rows = list(csv.reader(infile))
    if not rows or not rows[0] or rows[0][0] != "signature":
        raise RuntimeError(f"unexpected signature-scoring table format: {source}")

    rows[0][1:] = replace_labels(rows[0][1:], labels, "signature-scoring table")
    with destination.open("w", newline="", encoding="utf-8") as outfile:
        csv.writer(outfile, lineterminator="\n").writerows(rows)


def export_results(
    plan: dict[Path, Path],
    labels: dict[str, str],
    output_dir: Path,
    *,
    force: bool,
    dry_run: bool,
) -> None:
    output_dir = output_dir.resolve()
    existing = [output_dir / path for path in plan if (output_dir / path).exists()]
    if existing and not force:
        paths = ", ".join(str(path) for path in existing)
        raise FileExistsError(
            f"exported result already exists: {paths}; use --force to replace it"
        )

    print("scope: integrated")
    label_summary = ", ".join(f"{key}={value}" for key, value in labels.items())
    print(f"cluster labels: {label_summary}")
    print(f"output: {output_dir}")
    for relative_path, source in plan.items():
        print(f"{source} -> {output_dir / relative_path}")

    if dry_run:
        print("dry run: no files copied")
        return

    exporters = {
        DEA_RESULT: export_dea,
        GOEA_RESULT: export_goea,
        SCORING_RESULT: export_scoring,
    }
    for relative_path, source in plan.items():
        destination = output_dir / relative_path
        destination.parent.mkdir(parents=True, exist_ok=True)
        exporters[relative_path](source, destination, labels)

    print("export complete")


def main() -> None:
    args = parse_args()
    try:
        project_dir = args.project_dir.resolve()
        plan = build_export_plan(project_dir)
        labels = load_cluster_labels(project_dir)
        export_results(
            plan,
            labels,
            args.output_dir,
            force=args.force,
            dry_run=args.dry_run,
        )
    except (OSError, RuntimeError) as error:
        raise SystemExit(f"error: {error}") from None


if __name__ == "__main__":
    main()
